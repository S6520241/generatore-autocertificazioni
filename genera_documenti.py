"""
Generatore di autocertificazioni di ricevimento merce
(art. 45-bis Reg. UE 282/2011) a partire dalle fatture PDF.

Requisiti:
    pip install pdfplumber python-docx
    LibreOffice installato (serve per la conversione Word -> PDF)

I template devono stare nella stessa cartella di questo script:
    template_autocertificazione_it.docx
    template_autocertificazione_en.docx

Avvio con interfaccia grafica:
    python genera_documenti.py

Avvio da riga di comando (senza interfaccia):
    python genera_documenti.py --lingua en fattura1.pdf fattura2.pdf
"""

from __future__ import annotations

import argparse
import copy
import os
import queue
import re
import shutil
import subprocess
import sys
import tempfile
import threading
from dataclasses import dataclass, field
from datetime import date, timedelta
from pathlib import Path

import pdfplumber
from docx import Document
from docx.oxml.ns import qn
from docx.text.run import Run

import tkinter as tk
from tkinter import ttk, filedialog, messagebox


# ---------------------------------------------------------------------------
# Configurazione
# ---------------------------------------------------------------------------

if getattr(sys, 'frozen', False):
    # Programma .exe creato con PyInstaller
    CARTELLA_SCRIPT = Path(sys.executable).resolve().parent              # cartella dell'.exe
    CARTELLA_RISORSE = Path(getattr(sys, '_MEIPASS', CARTELLA_SCRIPT))   # file inclusi nell'.exe
else:
    CARTELLA_SCRIPT = Path(__file__).resolve().parent
    CARTELLA_RISORSE = CARTELLA_SCRIPT


def _percorso_template(nome: str) -> Path:
    """Un template messo accanto al programma ha la precedenza su quello incluso nell'.exe."""
    accanto = CARTELLA_SCRIPT / nome
    return accanto if accanto.exists() else CARTELLA_RISORSE / nome


TEMPLATES = {
    'it': _percorso_template('template_autocertificazione_it.docx'),
    'en': _percorso_template('template_autocertificazione_en.docx'),
}
NOMI_LINGUE = {'it': 'Italiano', 'en': 'Inglese'}

CARTELLA_OUTPUT_DEFAULT = CARTELLA_SCRIPT / 'DocumentiOutput'

FORMATO_DATA = '%d/%m/%Y'      # formato delle date scritte nell'autocertificazione
GIORNI_RICEVIMENTO = 2          # DATA_RICEVIMENTO = DATA_FATTURA + 2 giorni
GIORNI_FIRMA = 2                # DATA_FIRMA = DATA_RICEVIMENTO + 2 giorni

PERCORSI_LIBREOFFICE = [
    r'C:\Program Files\LibreOffice\program\soffice.exe',
    r'C:\Program Files (x86)\LibreOffice\program\soffice.exe',
    '/Applications/LibreOffice.app/Contents/MacOS/soffice',
]


class ErroreFattura(Exception):
    """Dati non trovati o non leggibili nella fattura PDF."""


class ErroreGenerazione(Exception):
    """Problema nella creazione del documento o nella conversione in PDF."""


# ---------------------------------------------------------------------------
# Dati della fattura
# ---------------------------------------------------------------------------

def formatta_data(d: date | None) -> str:
    return d.strftime(FORMATO_DATA) if d else ''


@dataclass
class DatiFattura:
    file_pdf: Path
    nome_cliente: str = ''
    indirizzo_cliente: str = ''
    prodotti: list = field(default_factory=list)
    numero_fattura: str = ''
    data_fattura: date | None = None

    @property
    def data_ricevimento(self) -> date | None:
        if not self.data_fattura:
            return None
        return self.data_fattura + timedelta(days=GIORNI_RICEVIMENTO)

    @property
    def data_firma(self) -> date | None:
        if not self.data_ricevimento:
            return None
        return self.data_ricevimento + timedelta(days=GIORNI_FIRMA)

    def campi_mancanti(self) -> list:
        mancanti = []
        if not self.nome_cliente:
            mancanti.append('nome cliente')
        if not self.indirizzo_cliente:
            mancanti.append('indirizzo')
        if not self.prodotti:
            mancanti.append('prodotti')
        if not self.numero_fattura:
            mancanti.append('numero fattura')
        if not self.data_fattura:
            mancanti.append('data fattura')
        return mancanti

    def valori_segnaposto(self) -> dict:
        return {
            'NOME_CLIENTE': self.nome_cliente,
            'INDIRIZZO_CLIENTE': self.indirizzo_cliente,
            'DESCRIZIONE_BENI': '\n'.join(self.prodotti),
            'NUMERO_FATTURA': self.numero_fattura,
            'DATA_FATTURA': formatta_data(self.data_fattura),
            'LUOGO_ARRIVO': self.indirizzo_cliente,
            'DATA_RICEVIMENTO': formatta_data(self.data_ricevimento),
            'DATA_FIRMA': formatta_data(self.data_firma),
            'PERSONA_ACCETTAZIONE': '',
        }


# ---------------------------------------------------------------------------
# Lettura della fattura PDF
# ---------------------------------------------------------------------------

# "nr. FPR 215/24 del 29/06/2024"  ->  numero "215", data 29/06/2024
RE_NUMERO = re.compile(r'\bnr\.?\s*FPR\s*([^/\s]+)\s*/', re.IGNORECASE)
RE_DATA = re.compile(
    r'\bnr\.?\s*FPR\s*\S+\s+del\s+(\d{1,2})[/.\-](\d{1,2})[/.\-](\d{2,4})',
    re.IGNORECASE,
)

# Righe del riquadro CLIENTE con l'identificativo fiscale ("Id Fiscale: ...")
RE_ID_FISCALE = re.compile(
    r'^(Id\s*Fiscale|P\.?\s*IVA|Partita\s*IVA|C\.?\s*F\.?|Codice\s*Fiscale)\s*:',
    re.IGNORECASE,
)
# Righe che chiudono l'indirizzo del cliente
RE_FINE_INDIRIZZO = re.compile(
    r'^(Codice\s*destinatario|PEC|Tel|Telefono|Cell|Fax|E-?mail)\b',
    re.IGNORECASE,
)
# Piè di pagina (per fatture su più pagine)
RE_PIE_PAGINA = re.compile(r'^(Copia analogica|Fattura\s+Nr\.?)', re.IGNORECASE)


def _raggruppa_righe(parole: list, tolleranza: float = 3.0) -> list:
    """Raggruppa le parole estratte da pdfplumber in righe (per coordinata verticale)."""
    righe = []
    for w in sorted(parole, key=lambda w: (w['top'], w['x0'])):
        if righe and abs(w['top'] - righe[-1][0]['top']) <= tolleranza:
            righe[-1].append(w)
        else:
            righe.append([w])
    return [sorted(r, key=lambda w: w['x0']) for r in righe]


def _testo(riga: list) -> str:
    return ' '.join(w['text'] for w in riga)


def _estrai_cliente(pagina) -> tuple:
    """Nome e indirizzo dal riquadro CLIENTE in alto a destra."""
    parole = pagina.extract_words(extra_attrs=['size'])
    intestazione = next((w for w in parole if w['text'].upper() == 'CLIENTE'), None)
    if intestazione is None:
        raise ErroreFattura('riquadro "CLIENTE" non trovato')

    # Il riquadro finisce dove inizia la sezione "PRODOTTI E SERVIZI"
    limite = min(
        (w['top'] for w in parole
         if w['text'].upper() == 'PRODOTTI' and w['top'] > intestazione['bottom']),
        default=pagina.height,
    )
    zona = [
        w for w in parole
        if w['x0'] >= intestazione['x0'] - 5 and intestazione['bottom'] < w['top'] < limite
    ]
    righe = [_testo(r) for r in _raggruppa_righe(zona)]

    indici_fiscali = [i for i, r in enumerate(righe) if RE_ID_FISCALE.match(r)]
    if not indici_fiscali:
        raise ErroreFattura('riga "Id Fiscale" del cliente non trovata')
    primo = indici_fiscali[0]

    # Nome: tutte le righe tra "CLIENTE" e "Id Fiscale" (può andare a capo)
    nome = ' '.join(righe[:primo]).strip()

    # Indirizzo: le righe sotto "Id Fiscale" fino a "Codice destinatario"
    indirizzo = []
    for r in righe[primo + 1:]:
        if RE_FINE_INDIRIZZO.match(r):
            break
        if RE_ID_FISCALE.match(r):
            continue
        indirizzo.append(r)

    return nome, ', '.join(indirizzo).strip()


def _estrai_prodotti(pdf) -> list:
    """Descrizioni della tabella "PRODOTTI E SERVIZI" (una voce per prodotto)."""
    prodotti = []
    for pagina in pdf.pages:
        righe = _raggruppa_righe(pagina.extract_words(extra_attrs=['size']))

        # Intestazione della tabella: riga con DESCRIZIONE e QUANTITA'
        for i, riga in enumerate(righe):
            col_descr = next((w for w in riga if w['text'].upper() == 'DESCRIZIONE'), None)
            col_qta = next((w for w in riga if w['text'].upper().startswith('QUANTIT')), None)
            if col_descr and col_qta:
                break
        else:
            continue  # nessuna tabella prodotti in questa pagina

        dim_intestazione = max(w['size'] for w in riga)
        x_descr = col_descr['x0'] - 2      # inizio colonna DESCRIZIONE
        x_qta = col_qta['x0'] - 1          # inizio colonna QUANTITA'

        for riga in righe[i + 1:]:
            # La tabella finisce al titolo della sezione successiva
            # ("METODO DI PAGAMENTO", scritto più grande) o al piè di pagina
            if max(w['size'] for w in riga) >= dim_intestazione + 2:
                break
            if RE_PIE_PAGINA.match(_testo(riga)):
                break

            numero = [w for w in riga if w['x1'] <= x_descr and w['text'].isdigit()]
            descrizione = ' '.join(
                w['text'] for w in riga if x_descr <= w['x0'] < x_qta
            ).strip()
            if not descrizione:
                continue
            if numero or not prodotti:
                prodotti.append(descrizione)               # nuovo prodotto
            else:
                prodotti[-1] += ' ' + descrizione          # descrizione su più righe
    return prodotti


def estrai_dati_fattura(pdf_path) -> DatiFattura:
    """Legge una fattura PDF e restituisce i dati per l'autocertificazione."""
    pdf_path = Path(pdf_path)
    dati = DatiFattura(file_pdf=pdf_path)
    try:
        pdf = pdfplumber.open(pdf_path)
    except Exception as e:
        raise ErroreFattura(f'impossibile aprire il PDF ({e})')

    with pdf:
        if not pdf.pages:
            raise ErroreFattura('il PDF non ha pagine')
        prima = pdf.pages[0]
        testo = prima.extract_text() or ''
        if not testo.strip():
            raise ErroreFattura('nessun testo nel PDF (è una scansione?)')

        m = RE_NUMERO.search(testo)
        if not m:
            raise ErroreFattura('numero fattura "nr. FPR .../.." non trovato')
        dati.numero_fattura = m.group(1)

        m = RE_DATA.search(testo)
        if not m:
            raise ErroreFattura('data fattura non trovata')
        giorno, mese, anno = (int(x) for x in m.groups())
        if anno < 100:
            anno += 2000
        try:
            dati.data_fattura = date(anno, mese, giorno)
        except ValueError:
            raise ErroreFattura(f'data fattura non valida: {m.group(0)}')

        dati.nome_cliente, dati.indirizzo_cliente = _estrai_cliente(prima)
        dati.prodotti = _estrai_prodotti(pdf)

    mancanti = dati.campi_mancanti()
    if mancanti:
        raise ErroreFattura('dati non trovati: ' + ', '.join(mancanti))
    return dati


# ---------------------------------------------------------------------------
# Compilazione del template Word
# ---------------------------------------------------------------------------

RE_SEGNAPOSTO = re.compile(r'\{\{\s*([A-Z_]+)\s*\}\}')


def sostituisci_nel_paragrafo(paragrafo, valori: dict) -> None:
    """
    Sostituisce i segnaposto {{...}} nel paragrafo mantenendo la formattazione
    del template (corsivo, dimensione, ecc.). I segnaposto possono essere spezzati
    su più "run" di Word. I valori inseriti sono in grassetto e sottolineati;
    gli "a capo" (\\n) diventano interruzioni di riga.
    """
    runs = list(paragrafo.runs)
    if not runs:
        return
    testi = [r.text for r in runs]
    completo = ''.join(testi)
    if '{' not in completo:
        return
    trovati = [m for m in RE_SEGNAPOSTO.finditer(completo) if m.group(1) in valori]
    if not trovati:
        return

    # A quale run appartiene ogni carattere del testo del paragrafo
    proprietario = []
    for indice, t in enumerate(testi):
        proprietario.extend([indice] * len(t))

    segmenti = []  # (testo, indice run da cui copiare lo stile, è_un_valore)

    def aggiungi_testo_originale(inizio, fine):
        s = inizio
        while s < fine:
            indice = proprietario[s]
            e = s
            while e < fine and proprietario[e] == indice:
                e += 1
            segmenti.append((completo[s:e], indice, False))
            s = e

    posizione = 0
    for m in trovati:
        aggiungi_testo_originale(posizione, m.start())
        valore = str(valori[m.group(1)])
        # Elenco su più righe dopo del testo (es. "the following goods"):
        # si parte da una riga nuova così l'elenco resta allineato
        if '\n' in valore and completo[:m.start()].strip():
            valore = '\n' + valore
        segmenti.append((valore, proprietario[m.start()], True))
        posizione = m.end()
    aggiungi_testo_originale(posizione, len(completo))

    primo_run = runs[0]._r
    for testo, indice, e_valore in segmenti:
        if not testo:
            continue
        nuovo = copy.deepcopy(runs[indice]._r)
        primo_run.addprevious(nuovo)
        run = Run(nuovo, paragrafo)
        run.text = testo  # gestisce \n (a capo) e \t (tab)
        if e_valore:
            rpr = nuovo.get_or_add_rPr()
            # Toglie la spaziatura dei caratteri del segnaposto (es. "{ { L U O G O ...")
            for spaziatura in rpr.findall(qn('w:spacing')):
                rpr.remove(spaziatura)
            run.bold = True
            run.underline = True

    for r in runs:
        r._r.getparent().remove(r._r)


def _tutti_i_paragrafi(doc):
    """Tutti i paragrafi del documento: corpo, tabelle (anche annidate), intestazioni e piè di pagina."""
    def da_contenitore(contenitore):
        yield from contenitore.paragraphs
        for tabella in contenitore.tables:
            for riga in tabella.rows:
                for cella in riga.cells:
                    yield from da_contenitore(cella)

    yield from da_contenitore(doc)
    for sezione in doc.sections:
        for parte in (sezione.header, sezione.footer,
                      sezione.first_page_header, sezione.first_page_footer,
                      sezione.even_page_header, sezione.even_page_footer):
            if not parte.is_linked_to_previous:
                yield from da_contenitore(parte)


# ---------------------------------------------------------------------------
# Conversione in PDF e salvataggio
# ---------------------------------------------------------------------------

def trova_libreoffice() -> str | None:
    for percorso in PERCORSI_LIBREOFFICE:
        if Path(percorso).exists():
            return percorso
    for nome in ('soffice', 'libreoffice'):
        trovato = shutil.which(nome)
        if trovato:
            return trovato
    return None


def converti_in_pdf(docx_path: Path, output_dir: Path, soffice: str) -> Path:
    """Converte un .docx in PDF con LibreOffice e restituisce il percorso del PDF."""
    # Profilo LibreOffice dedicato: la conversione funziona anche se LibreOffice è già aperto
    profilo = Path(tempfile.gettempdir()) / 'autocertificazioni_libreoffice'
    comando = [
        soffice,
        f'-env:UserInstallation={profilo.as_uri()}',
        '--headless',
        '--convert-to', 'pdf',
        '--outdir', str(output_dir),
        str(docx_path),
    ]
    opzioni = {}
    if os.name == 'nt':
        opzioni['creationflags'] = subprocess.CREATE_NO_WINDOW  # niente finestra nera
    try:
        esito = subprocess.run(comando, stdout=subprocess.PIPE, stderr=subprocess.PIPE,
                               timeout=180, **opzioni)
    except subprocess.TimeoutExpired:
        raise ErroreGenerazione('LibreOffice non ha risposto entro 3 minuti')
    except OSError as e:
        raise ErroreGenerazione(f'impossibile avviare LibreOffice ({e})')

    pdf = output_dir / (docx_path.stem + '.pdf')
    if not pdf.exists():
        dettaglio = esito.stderr.decode(errors='ignore').strip()[:300]
        raise ErroreGenerazione('conversione in PDF fallita' + (f': {dettaglio}' if dettaglio else ''))
    return pdf


def nome_sicuro(testo: str) -> str:
    pulito = ''.join(c for c in testo if c.isalnum() or c in (' ', '_', '-'))
    pulito = re.sub(r'\s+', ' ', pulito).strip()
    return pulito or 'Cliente'


def genera_autocertificazione(dati: DatiFattura, lingua: str, cartella_output, soffice: str) -> Path:
    """Compila il template nella lingua scelta e salva il PDF nella cartella del cliente."""
    template = TEMPLATES[lingua]
    if not template.exists():
        raise ErroreGenerazione(f'template non trovato: {template.name}')
    mancanti = dati.campi_mancanti()
    if mancanti:
        raise ErroreGenerazione('dati mancanti: ' + ', '.join(mancanti))

    doc = Document(str(template))
    valori = dati.valori_segnaposto()
    for paragrafo in _tutti_i_paragrafi(doc):
        sostituisci_nel_paragrafo(paragrafo, valori)

    cartella_cliente = Path(cartella_output) / nome_sicuro(dati.nome_cliente)
    cartella_cliente.mkdir(parents=True, exist_ok=True)

    numero = dati.numero_fattura.replace('/', '_').replace('\\', '_')
    nome_file = f'Dichiarazione_Fattura_{numero}_data_{dati.data_fattura:%d-%m-%Y}'
    if lingua != 'it':
        nome_file += f'_{lingua.upper()}'
    destinazione = cartella_cliente / f'{nome_file}.pdf'

    cartella_temp = Path(tempfile.mkdtemp(prefix='autocert_'))
    try:
        docx_temp = cartella_temp / f'{nome_file}.docx'
        doc.save(str(docx_temp))
        pdf_temp = converti_in_pdf(docx_temp, cartella_temp, soffice)
        try:
            shutil.copyfile(pdf_temp, destinazione)
        except PermissionError:
            raise ErroreGenerazione(f'impossibile scrivere {destinazione.name}: il file è aperto in un altro programma?')
    finally:
        shutil.rmtree(cartella_temp, ignore_errors=True)
    return destinazione


# ---------------------------------------------------------------------------
# Interfaccia grafica
# ---------------------------------------------------------------------------

def apri_cartella(percorso: Path) -> None:
    percorso = Path(percorso)
    percorso.mkdir(parents=True, exist_ok=True)
    if os.name == 'nt':
        os.startfile(str(percorso))
    elif sys.platform == 'darwin':
        subprocess.Popen(['open', str(percorso)])
    else:
        subprocess.Popen(['xdg-open', str(percorso)])


class FinestraModifica(tk.Toplevel):
    """Finestra per correggere a mano i dati letti da una fattura."""

    def __init__(self, padre, dati: DatiFattura, al_salvataggio):
        super().__init__(padre)
        self.dati = dati
        self.al_salvataggio = al_salvataggio
        self.title(f'Modifica dati - {dati.file_pdf.name}')
        self.transient(padre)
        self.resizable(True, True)

        corpo = ttk.Frame(self, padding=14)
        corpo.pack(fill='both', expand=True)
        corpo.columnconfigure(1, weight=1)

        self.var_numero = tk.StringVar(value=dati.numero_fattura)
        self.var_data = tk.StringVar(value=formatta_data(dati.data_fattura))
        self.var_nome = tk.StringVar(value=dati.nome_cliente)
        self.var_indirizzo = tk.StringVar(value=dati.indirizzo_cliente)

        campi = [
            ('Numero fattura', self.var_numero),
            ('Data fattura (gg/mm/aaaa)', self.var_data),
            ('Nome cliente', self.var_nome),
            ('Indirizzo / luogo di arrivo', self.var_indirizzo),
        ]
        for riga, (etichetta, variabile) in enumerate(campi):
            ttk.Label(corpo, text=etichetta).grid(row=riga, column=0, sticky='w', pady=3, padx=(0, 10))
            ttk.Entry(corpo, textvariable=variabile, width=60).grid(row=riga, column=1, sticky='ew', pady=3)

        ttk.Label(corpo, text='Prodotti (uno per riga)').grid(row=4, column=0, sticky='nw', pady=3, padx=(0, 10))
        self.testo_prodotti = tk.Text(corpo, width=60, height=8, wrap='word', font='TkDefaultFont')
        self.testo_prodotti.grid(row=4, column=1, sticky='nsew', pady=3)
        self.testo_prodotti.insert('1.0', '\n'.join(dati.prodotti))
        corpo.rowconfigure(4, weight=1)

        pulsanti = ttk.Frame(corpo)
        pulsanti.grid(row=5, column=0, columnspan=2, sticky='e', pady=(10, 0))
        ttk.Button(pulsanti, text='Annulla', command=self.destroy).pack(side='right')
        ttk.Button(pulsanti, text='Salva', command=self._salva).pack(side='right', padx=(0, 6))

        self.bind('<Escape>', lambda e: self.destroy())
        self.grab_set()

    def _salva(self):
        testo_data = self.var_data.get().strip()
        nuova_data = None
        if testo_data:
            m = re.fullmatch(r'(\d{1,2})[/.\-](\d{1,2})[/.\-](\d{4})', testo_data)
            try:
                if not m:
                    raise ValueError
                nuova_data = date(int(m.group(3)), int(m.group(2)), int(m.group(1)))
            except ValueError:
                messagebox.showerror('Data non valida', 'Scrivi la data nel formato gg/mm/aaaa.', parent=self)
                return
        self.dati.numero_fattura = self.var_numero.get().strip()
        self.dati.data_fattura = nuova_data
        self.dati.nome_cliente = self.var_nome.get().strip()
        self.dati.indirizzo_cliente = self.var_indirizzo.get().strip()
        self.dati.prodotti = [
            r.strip() for r in self.testo_prodotti.get('1.0', 'end').splitlines() if r.strip()
        ]
        self.al_salvataggio(self.dati)
        self.destroy()


class App:
    COLONNE = ('file', 'numero', 'data', 'cliente', 'prodotti', 'stato')

    def __init__(self, root: tk.Tk):
        self.root = root
        root.title('Generatore autocertificazioni')
        root.geometry('1040x720')
        root.minsize(860, 600)

        self.lingua = tk.StringVar(value='it')
        self.cartella_output = tk.StringVar(value=str(CARTELLA_OUTPUT_DEFAULT))
        self.fatture = {}        # iid -> DatiFattura
        self.ultima_cartella = str(Path.home())
        self.coda = queue.Queue()
        self.in_esecuzione = False

        self._costruisci()
        self.root.after(100, self._leggi_coda)

    # -- costruzione finestra --------------------------------------------

    def _costruisci(self):
        principale = ttk.Frame(self.root, padding=12)
        principale.pack(fill='both', expand=True)
        principale.columnconfigure(0, weight=1)
        principale.rowconfigure(1, weight=1)
        principale.rowconfigure(2, weight=1)

        # 1. Lingua + cartella di destinazione
        alto = ttk.Frame(principale)
        alto.grid(row=0, column=0, sticky='ew')
        alto.columnconfigure(1, weight=1)

        box_lingua = ttk.LabelFrame(alto, text=' Lingua dell\'autocertificazione ', padding=(10, 6))
        box_lingua.grid(row=0, column=0, sticky='nsw', padx=(0, 10))
        ttk.Radiobutton(box_lingua, text='Italiano', value='it', variable=self.lingua).pack(side='left', padx=(0, 16))
        ttk.Radiobutton(box_lingua, text='Inglese (English)', value='en', variable=self.lingua).pack(side='left')

        box_output = ttk.LabelFrame(alto, text=' Cartella di destinazione ', padding=(10, 6))
        box_output.grid(row=0, column=1, sticky='nsew')
        box_output.columnconfigure(0, weight=1)
        ttk.Entry(box_output, textvariable=self.cartella_output).grid(row=0, column=0, sticky='ew', padx=(0, 6))
        ttk.Button(box_output, text='Sfoglia…', command=self._scegli_cartella).grid(row=0, column=1)
        ttk.Button(box_output, text='Apri', command=self._apri_output).grid(row=0, column=2, padx=(6, 0))

        # 2. Elenco fatture
        box_fatture = ttk.LabelFrame(principale, text=' Fatture PDF ', padding=(10, 6))
        box_fatture.grid(row=1, column=0, sticky='nsew', pady=(10, 0))
        box_fatture.columnconfigure(0, weight=1)
        box_fatture.rowconfigure(1, weight=1)

        barra = ttk.Frame(box_fatture)
        barra.grid(row=0, column=0, columnspan=2, sticky='ew', pady=(0, 6))
        self.btn_aggiungi = ttk.Button(barra, text='Carica fatture PDF…', command=self._aggiungi_fatture)
        self.btn_aggiungi.pack(side='left')
        self.btn_modifica = ttk.Button(barra, text='Modifica dati…', command=self._modifica_selezionata)
        self.btn_modifica.pack(side='left', padx=(6, 0))
        self.btn_rimuovi = ttk.Button(barra, text='Rimuovi selezionate', command=self._rimuovi_selezionate)
        self.btn_rimuovi.pack(side='left', padx=(6, 0))
        self.btn_svuota = ttk.Button(barra, text='Svuota elenco', command=self._svuota)
        self.btn_svuota.pack(side='left', padx=(6, 0))
        self.lbl_conteggio = ttk.Label(barra, text='')
        self.lbl_conteggio.pack(side='right')

        self.tabella = ttk.Treeview(box_fatture, columns=self.COLONNE, show='headings', selectmode='extended')
        intestazioni = {
            'file': ('File', 250), 'numero': ('N. fattura', 80), 'data': ('Data', 100),
            'cliente': ('Cliente', 250), 'prodotti': ('Prodotti', 70), 'stato': ('Stato', 200),
        }
        for colonna, (titolo, larghezza) in intestazioni.items():
            self.tabella.heading(colonna, text=titolo, anchor='w')
            self.tabella.column(colonna, width=larghezza, anchor='w',
                                stretch=colonna in ('file', 'cliente', 'stato'))
        self.tabella.column('prodotti', anchor='center')
        self.tabella.tag_configure('errore', foreground='#b3261e')
        self.tabella.tag_configure('fatto', foreground='#1e7b34')
        self.tabella.grid(row=1, column=0, sticky='nsew')
        scroll = ttk.Scrollbar(box_fatture, orient='vertical', command=self.tabella.yview)
        scroll.grid(row=1, column=1, sticky='ns')
        self.tabella.configure(yscrollcommand=scroll.set)
        self.tabella.bind('<<TreeviewSelect>>', lambda e: self._mostra_dettaglio())
        self.tabella.bind('<Double-1>', lambda e: self._modifica_selezionata())
        self.tabella.bind('<Delete>', lambda e: self._rimuovi_selezionate())

        # 3. Dettaglio + registro
        basso = ttk.Frame(principale)
        basso.grid(row=2, column=0, sticky='nsew', pady=(10, 0))
        basso.columnconfigure(0, weight=1)
        basso.columnconfigure(1, weight=1)
        basso.rowconfigure(0, weight=1)

        box_dettaglio = ttk.LabelFrame(basso, text=' Dati che verranno inseriti ', padding=(8, 6))
        box_dettaglio.grid(row=0, column=0, sticky='nsew', padx=(0, 10))
        box_dettaglio.columnconfigure(0, weight=1)
        box_dettaglio.rowconfigure(0, weight=1)
        self.dettaglio = tk.Text(box_dettaglio, height=10, wrap='word', relief='flat',
                                 font='TkDefaultFont', background=self.root.cget('background'))
        self.dettaglio.grid(row=0, column=0, sticky='nsew')
        scroll_det = ttk.Scrollbar(box_dettaglio, orient='vertical', command=self.dettaglio.yview)
        scroll_det.grid(row=0, column=1, sticky='ns')
        self.dettaglio.configure(yscrollcommand=scroll_det.set)
        self.dettaglio.tag_configure('etichetta', foreground='#555555')
        self._scrivi_dettaglio('Seleziona una fattura per vedere i dati letti dal PDF.\n'
                               'Doppio clic su una riga per correggerli.')

        box_registro = ttk.LabelFrame(basso, text=' Registro ', padding=(8, 6))
        box_registro.grid(row=0, column=1, sticky='nsew')
        box_registro.columnconfigure(0, weight=1)
        box_registro.rowconfigure(0, weight=1)
        self.registro = tk.Text(box_registro, height=10, wrap='word', state='disabled', font='TkDefaultFont')
        self.registro.grid(row=0, column=0, sticky='nsew')
        scroll_reg = ttk.Scrollbar(box_registro, orient='vertical', command=self.registro.yview)
        scroll_reg.grid(row=0, column=1, sticky='ns')
        self.registro.configure(yscrollcommand=scroll_reg.set)
        self.registro.tag_configure('errore', foreground='#b3261e')
        self.registro.tag_configure('ok', foreground='#1e7b34')

        # 4. Avanzamento + pulsante principale
        fondo = ttk.Frame(principale)
        fondo.grid(row=3, column=0, sticky='ew', pady=(10, 0))
        fondo.columnconfigure(0, weight=1)
        self.avanzamento = ttk.Progressbar(fondo, mode='determinate')
        self.avanzamento.grid(row=0, column=0, sticky='ew', padx=(0, 10))
        self.btn_genera = ttk.Button(fondo, text='Genera autocertificazioni', command=self._genera)
        self.btn_genera.grid(row=0, column=1, ipadx=10, ipady=4)

        self._aggiorna_conteggio()

    # -- utilità ---------------------------------------------------------

    def _log(self, messaggio: str, tag: str | None = None):
        self.registro.configure(state='normal')
        self.registro.insert('end', messaggio + '\n', tag or ())
        self.registro.see('end')
        self.registro.configure(state='disabled')

    def _scrivi_dettaglio(self, contenuto):
        """contenuto: stringa semplice oppure lista di (etichetta, valore)."""
        self.dettaglio.configure(state='normal')
        self.dettaglio.delete('1.0', 'end')
        if isinstance(contenuto, str):
            self.dettaglio.insert('end', contenuto, 'etichetta')
        else:
            for etichetta, valore in contenuto:
                self.dettaglio.insert('end', f'{etichetta}: ', 'etichetta')
                self.dettaglio.insert('end', f'{valore}\n')
        self.dettaglio.configure(state='disabled')

    def _aggiorna_conteggio(self):
        n = len(self.fatture)
        pronte = sum(1 for d in self.fatture.values() if not d.campi_mancanti())
        self.lbl_conteggio.configure(text=f'{n} fatture caricate, {pronte} pronte' if n else 'Nessuna fattura caricata')

    def _aggiorna_riga(self, iid: str, stato: str | None = None, tag: str | None = None):
        dati = self.fatture[iid]
        mancanti = dati.campi_mancanti()
        if stato is None:
            stato = 'Pronta' if not mancanti else 'Da completare: ' + ', '.join(mancanti)
        if tag is None:
            tag = 'errore' if mancanti else ''
        self.tabella.item(iid, values=(
            dati.file_pdf.name, dati.numero_fattura, formatta_data(dati.data_fattura),
            dati.nome_cliente, len(dati.prodotti) or '', stato,
        ), tags=(tag,) if tag else ())

    # -- azioni ----------------------------------------------------------

    def _aggiungi_fatture(self):
        percorsi = filedialog.askopenfilenames(
            parent=self.root, title='Seleziona le fatture PDF', initialdir=self.ultima_cartella,
            filetypes=[('Fatture PDF', '*.pdf'), ('Tutti i file', '*.*')],
        )
        if not percorsi:
            return
        self.ultima_cartella = str(Path(percorsi[0]).parent)
        gia_presenti = {d.file_pdf.resolve() for d in self.fatture.values()}

        self.root.configure(cursor='watch')
        self.root.update_idletasks()
        try:
            for percorso in percorsi:
                percorso = Path(percorso)
                if percorso.resolve() in gia_presenti:
                    continue
                iid = self.tabella.insert('', 'end', values=(percorso.name,))
                try:
                    dati = estrai_dati_fattura(percorso)
                    self.fatture[iid] = dati
                    self._aggiorna_riga(iid)
                    self._log(f'Letta {percorso.name}: fattura {dati.numero_fattura} del '
                              f'{formatta_data(dati.data_fattura)}, {len(dati.prodotti)} prodotti')
                except Exception as e:
                    # Si tiene la fattura in elenco: i dati si possono completare a mano
                    self.fatture[iid] = DatiFattura(file_pdf=percorso)
                    self._aggiorna_riga(iid, stato=f'Errore: {e}', tag='errore')
                    self._log(f'{percorso.name}: {e} (doppio clic per inserire i dati a mano)', 'errore')
                self.root.update_idletasks()
        finally:
            self.root.configure(cursor='')
        self._aggiorna_conteggio()
        figli = self.tabella.get_children()
        if figli and not self.tabella.selection():
            self.tabella.selection_set(figli[0])

    def _mostra_dettaglio(self):
        selezione = self.tabella.selection()
        if not selezione:
            return
        dati = self.fatture.get(selezione[0])
        if dati is None:
            return
        prodotti = ''.join(f'\n   • {p}' for p in dati.prodotti) or ' —'
        self._scrivi_dettaglio([
            ('File', dati.file_pdf.name),
            ('Cliente', dati.nome_cliente or '—'),
            ('Indirizzo / luogo di arrivo', dati.indirizzo_cliente or '—'),
            ('Numero fattura', dati.numero_fattura or '—'),
            ('Data fattura', formatta_data(dati.data_fattura) or '—'),
            ('Data ricevimento', formatta_data(dati.data_ricevimento) or '—'),
            ('Data firma', formatta_data(dati.data_firma) or '—'),
            (f'Prodotti ({len(dati.prodotti)})', prodotti),
        ])

    def _modifica_selezionata(self):
        if self.in_esecuzione:
            return
        selezione = self.tabella.selection()
        if not selezione:
            messagebox.showinfo('Modifica dati', 'Seleziona prima una fattura dall\'elenco.', parent=self.root)
            return
        iid = selezione[0]

        def salvato(_dati):
            self._aggiorna_riga(iid)
            self._aggiorna_conteggio()
            self._mostra_dettaglio()
            self._log(f'Dati modificati a mano: {_dati.file_pdf.name}')

        FinestraModifica(self.root, self.fatture[iid], salvato)

    def _rimuovi_selezionate(self):
        if self.in_esecuzione:
            return
        for iid in self.tabella.selection():
            self.tabella.delete(iid)
            self.fatture.pop(iid, None)
        self._aggiorna_conteggio()
        self._scrivi_dettaglio('Seleziona una fattura per vedere i dati letti dal PDF.')

    def _svuota(self):
        if self.in_esecuzione:
            return
        self.tabella.delete(*self.tabella.get_children())
        self.fatture.clear()
        self._aggiorna_conteggio()
        self._scrivi_dettaglio('Seleziona una fattura per vedere i dati letti dal PDF.')

    def _scegli_cartella(self):
        cartella = filedialog.askdirectory(parent=self.root, title='Cartella di destinazione',
                                           initialdir=self.cartella_output.get() or str(Path.home()))
        if cartella:
            self.cartella_output.set(cartella)

    def _apri_output(self):
        try:
            apri_cartella(Path(self.cartella_output.get()))
        except Exception as e:
            messagebox.showerror('Errore', f'Impossibile aprire la cartella:\n{e}', parent=self.root)

    def _imposta_controlli(self, attivi: bool):
        stato = '!disabled' if attivi else 'disabled'
        for pulsante in (self.btn_aggiungi, self.btn_modifica, self.btn_rimuovi,
                         self.btn_svuota, self.btn_genera):
            pulsante.state([stato])

    def _genera(self):
        if self.in_esecuzione:
            return
        if not self.fatture:
            messagebox.showwarning('Nessuna fattura', 'Carica almeno una fattura PDF.', parent=self.root)
            return

        lingua = self.lingua.get()
        template = TEMPLATES[lingua]
        if not template.exists():
            messagebox.showerror('Template mancante',
                                 f'Non trovo il template:\n{template}\n\n'
                                 'Mettilo nella stessa cartella dello script.', parent=self.root)
            return

        soffice = trova_libreoffice()
        if not soffice:
            messagebox.showerror('LibreOffice non trovato',
                                 'Serve LibreOffice per creare i PDF.\n'
                                 'Installalo da https://www.libreoffice.org e riprova.', parent=self.root)
            return

        cartella = self.cartella_output.get().strip()
        if not cartella:
            messagebox.showwarning('Cartella', 'Scegli la cartella di destinazione.', parent=self.root)
            return

        da_fare = [(iid, d) for iid, d in self.fatture.items() if not d.campi_mancanti()]
        incomplete = len(self.fatture) - len(da_fare)
        if not da_fare:
            messagebox.showwarning('Dati incompleti',
                                   'Nessuna fattura ha tutti i dati necessari.\n'
                                   'Fai doppio clic su una riga per completarla.', parent=self.root)
            return
        if incomplete and not messagebox.askyesno(
                'Dati incompleti',
                f'{incomplete} fatture hanno dati mancanti e verranno saltate.\nContinuare con le altre?',
                parent=self.root):
            return

        self.in_esecuzione = True
        self._imposta_controlli(False)
        self.avanzamento.configure(maximum=len(da_fare), value=0)
        self._log(f'Generazione di {len(da_fare)} autocertificazioni in {NOMI_LINGUE[lingua].lower()}…')
        for iid, _ in da_fare:
            self._aggiorna_riga(iid, stato='In attesa…', tag='')

        threading.Thread(target=self._lavora, args=(da_fare, lingua, cartella, soffice), daemon=True).start()

    def _lavora(self, da_fare, lingua, cartella, soffice):
        """Eseguito in un thread separato per non bloccare la finestra."""
        riusciti = 0
        for n, (iid, dati) in enumerate(da_fare, start=1):
            self.coda.put(('stato', iid, 'Generazione…', ''))
            try:
                pdf = genera_autocertificazione(dati, lingua, cartella, soffice)
                riusciti += 1
                self.coda.put(('ok', iid, pdf))
            except Exception as e:
                self.coda.put(('errore', iid, str(e)))
            self.coda.put(('progresso', n))
        self.coda.put(('fine', riusciti, len(da_fare) - riusciti, cartella))

    def _leggi_coda(self):
        try:
            while True:
                messaggio = self.coda.get_nowait()
                tipo = messaggio[0]
                if tipo == 'stato':
                    _, iid, testo, tag = messaggio
                    if self.tabella.exists(iid):
                        self._aggiorna_riga(iid, stato=testo, tag=tag)
                elif tipo == 'ok':
                    _, iid, pdf = messaggio
                    if self.tabella.exists(iid):
                        self._aggiorna_riga(iid, stato=f'Creato ✓ {pdf.name}', tag='fatto')
                    self._log(f'[OK] {pdf.parent.name}/{pdf.name}', 'ok')
                elif tipo == 'errore':
                    _, iid, errore = messaggio
                    nome = self.fatture[iid].file_pdf.name if iid in self.fatture else ''
                    if self.tabella.exists(iid):
                        self._aggiorna_riga(iid, stato=f'Errore: {errore}', tag='errore')
                    self._log(f'[ERRORE] {nome}: {errore}', 'errore')
                elif tipo == 'progresso':
                    self.avanzamento.configure(value=messaggio[1])
                elif tipo == 'fine':
                    _, riusciti, falliti, cartella = messaggio
                    self.in_esecuzione = False
                    self._imposta_controlli(True)
                    self._log(f'Completato: {riusciti} creati, {falliti} errori.', 'ok' if not falliti else 'errore')
                    testo = f'Autocertificazioni create: {riusciti}'
                    if falliti:
                        testo += f'\nErrori: {falliti} (vedi registro)'
                    if riusciti and messagebox.askyesno('Fatto', testo + '\n\nAprire la cartella di destinazione?',
                                                        parent=self.root):
                        self._apri_output()
                    elif not riusciti:
                        messagebox.showerror('Errore', testo, parent=self.root)
        except queue.Empty:
            pass
        self.root.after(100, self._leggi_coda)


def avvia_interfaccia():
    if os.name == 'nt':
        try:  # testo nitido sugli schermi ad alta risoluzione
            import ctypes
            ctypes.windll.shcore.SetProcessDpiAwareness(1)
        except Exception:
            pass
    root = tk.Tk()
    stile = ttk.Style(root)
    if os.name != 'nt' and 'clam' in stile.theme_names():
        stile.theme_use('clam')
    App(root)
    root.mainloop()


# ---------------------------------------------------------------------------
# Riga di comando (facoltativa)
# ---------------------------------------------------------------------------

def main_riga_di_comando(argomenti):
    soffice = trova_libreoffice()
    if not soffice:
        print('ERRORE: LibreOffice non trovato.')
        return 1
    errori = 0
    for percorso in argomenti.fatture:
        try:
            dati = estrai_dati_fattura(percorso)
            pdf = genera_autocertificazione(dati, argomenti.lingua, argomenti.output, soffice)
            print(f'[OK] {Path(percorso).name} -> {pdf}')
        except Exception as e:
            errori += 1
            print(f'[ERRORE] {Path(percorso).name}: {e}')
    return 1 if errori else 0


if __name__ == '__main__':
    parser = argparse.ArgumentParser(description='Genera autocertificazioni dalle fatture PDF.')
    parser.add_argument('fatture', nargs='*', help='fatture PDF (senza argomenti si apre la finestra)')
    parser.add_argument('--lingua', choices=sorted(TEMPLATES), default='it')
    parser.add_argument('--output', default=str(CARTELLA_OUTPUT_DEFAULT))
    argomenti = parser.parse_args()

    if argomenti.fatture:
        sys.exit(main_riga_di_comando(argomenti))
    avvia_interfaccia()
